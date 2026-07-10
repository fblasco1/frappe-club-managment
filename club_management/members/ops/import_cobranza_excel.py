"""Importa cobros desde Excel «INFORME DE COBRANZAS» (formato legacy).

Uso:

    bench --site dev.localhost execute club_management.members.ops.import_cobranza_excel.run \\
        --kwargs '{"file_path": "/workspace/scripts/COBRANZA_01_AL_07-07.xlsx", "dry_run": true}'

    bench --site dev.localhost execute club_management.members.ops.import_cobranza_excel.run \\
        --kwargs '{"file_path": "/workspace/scripts/COBRANZA_01_AL_07-07.xlsx", "apply": true}'

Columnas esperadas por bloque de concepto:
  Fecha | Socio (número) | Nombre | Período | Importe | Cobrador | Comisión | | Notas
"""

from __future__ import annotations

import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

import frappe
from frappe.utils import flt, getdate, today

from club_management.integrations.payment_ledger_postgres import apply_patch
from club_management.members.services.cobranza_manual import (
	SALES_INVOICE_DOCTYPE,
	_campo_socio_en,
	format_periodo_cobro,
	registrar_cobro_manual,
)

MESES = {
	"ENERO": 1,
	"FEBRERO": 2,
	"MARZO": 3,
	"ABRIL": 4,
	"MAYO": 5,
	"JUNIO": 6,
	"JULIO": 7,
	"AGOSTO": 8,
	"SEPTIEMBRE": 9,
	"OCTUBRE": 10,
	"NOVIEMBRE": 11,
	"DICIEMBRE": 12,
}

CUOTA_KEYWORDS = ("CUOTA SOCIAL", "CARNET")


def _parse_periodo(raw: str | None, fecha: date) -> str | None:
	if not raw:
		return None
	text = str(raw).strip().upper()
	year_match = re.search(r"(20\d{2})", text)
	year = int(year_match.group(1)) if year_match else fecha.year
	for nombre, mes in MESES.items():
		if nombre in text:
			return f"{mes:02d}/{year}"
	if text in MESES:
		return f"{mes:02d}/{year}" if (mes := MESES[text]) else None
	return None


def _resolve_mode_of_payment(notas: str | None, cobrador: int | None) -> str:
	text = (notas or "").lower()
	if any(token in text for token in ("trans", "tran", "transfer")):
		return "Wire Transfer"
	if cobrador in (2, 3):
		return "Credit Card"
	return "Cash"


def _parse_excel_rows(file_path: str) -> list[dict[str, Any]]:
	from openpyxl import load_workbook

	wb = load_workbook(file_path, read_only=True, data_only=True)
	ws = wb[wb.sheetnames[0]]
	rows = list(ws.iter_rows(values_only=True))
	current_concept: str | None = None
	parsed: list[dict[str, Any]] = []

	for row in rows:
		if row[0] and str(row[0]).startswith("Concepto:"):
			current_concept = str(row[0]).replace("Concepto:", "").strip()
			continue
		if not row[0] or not hasattr(row[0], "year") or not row[1]:
			continue
		fecha_val = row[0].date() if hasattr(row[0], "date") else getdate(row[0])
		parsed.append(
			{
				"fecha": fecha_val,
				"socio": str(row[1]).strip(),
				"nombre": row[2],
				"periodo_raw": row[3],
				"importe": flt(row[4]),
				"cobrador": row[5],
				"notas": row[8] if len(row) > 8 else None,
				"concepto": current_concept,
			}
		)
	return parsed


def _pending_invoices(socio_name: str, periodo: str) -> list[dict[str, Any]]:
	campo_socio = _campo_socio_en(SALES_INVOICE_DOCTYPE)
	if not campo_socio:
		return []
	filters: dict[str, Any] = {
		campo_socio: socio_name,
		"docstatus": 1,
		"outstanding_amount": [">", 0],
	}
	if frappe.get_meta(SALES_INVOICE_DOCTYPE).has_field("periodo_cobro"):
		filters["periodo_cobro"] = periodo
	return frappe.get_all(
		SALES_INVOICE_DOCTYPE,
		filters=filters,
		fields=["name", "outstanding_amount", "grand_total", "remarks"],
		order_by="posting_date asc",
	)


def _invoice_line_amounts(invoice_name: str) -> list[float]:
	return [
		flt(row.amount)
		for row in frappe.get_all(
			"Sales Invoice Item",
			filters={"parent": invoice_name},
			fields=["amount"],
		)
	]


def _match_invoice(
	invoices: list[dict[str, Any]],
	importe: float,
	concepto: str | None,
) -> str | None:
	if not invoices:
		return None

	concepto_upper = (concepto or "").upper()
	is_cuota = any(k in concepto_upper for k in CUOTA_KEYWORDS)

	def _remarks_kind(inv: dict[str, Any]) -> str:
		remarks = (inv.get("remarks") or "").lower()
		if "aranceles" in remarks:
			return "arancel"
		if "cuota mensual" in remarks:
			return "cuota"
		return "otro"

	typed: list[dict[str, Any]] = []
	for inv in invoices:
		kind = _remarks_kind(inv)
		if is_cuota and kind == "cuota":
			typed.append(inv)
		elif not is_cuota and kind == "arancel":
			typed.append(inv)
	candidates = typed or invoices

	for inv in candidates:
		if flt(inv.outstanding_amount) == importe:
			return inv["name"]

	for inv in candidates:
		line_amounts = _invoice_line_amounts(inv["name"])
		if importe in line_amounts and flt(inv.outstanding_amount) == importe:
			return inv["name"]

	for inv in candidates:
		if len(candidates) == 1 and flt(inv.outstanding_amount) >= importe:
			return inv["name"]

	for inv in invoices:
		if flt(inv.outstanding_amount) == importe:
			return inv["name"]

	return None


def _serialize_row(row: dict[str, Any]) -> dict[str, Any]:
	out: dict[str, Any] = {}
	for key, value in row.items():
		if isinstance(value, date):
			out[key] = value.isoformat()
		else:
			out[key] = value
	return out


def _write_log(log_path: str, payload: dict[str, Any]) -> str:
	path = Path(log_path)
	path.parent.mkdir(parents=True, exist_ok=True)
	path.write_text(json.dumps(payload, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
	return str(path)


def run(
	file_path: str,
	dry_run: bool = True,
	apply: bool = False,
	fecha_desde: str | None = None,
	fecha_hasta: str | None = None,
	log_path: str | None = None,
) -> dict[str, Any]:
	"""Importa cobros del Excel; por defecto dry_run."""
	frappe.only_for(("System Manager", "Secretaria"))
	apply_patch()

	if apply:
		dry_run = False

	rows = _parse_excel_rows(file_path)
	desde = getdate(fecha_desde) if fecha_desde else None
	hasta = getdate(fecha_hasta) if fecha_hasta else None

	ok: list[dict[str, Any]] = []
	omitidos: list[dict[str, Any]] = []
	errores: list[dict[str, Any]] = []

	for row in rows:
		if desde and row["fecha"] < desde:
			continue
		if hasta and row["fecha"] > hasta:
			continue

		socio_name = row["socio"]
		if not frappe.db.exists("Socio", socio_name):
			errores.append({**row, "error": "Socio no encontrado"})
			continue

		periodo = _parse_periodo(row["periodo_raw"], row["fecha"])
		if not periodo:
			errores.append({**row, "error": f"Período no interpretable: {row['periodo_raw']}"})
			continue

		invoices = _pending_invoices(socio_name, periodo)
		invoice_name = _match_invoice(invoices, row["importe"], row["concepto"])
		if not invoice_name:
			omitidos.append({**row, "periodo": periodo, "motivo": "Sin factura pendiente coincidente"})
			continue

		outstanding = flt(frappe.db.get_value(SALES_INVOICE_DOCTYPE, invoice_name, "outstanding_amount"))
		if outstanding <= 0:
			omitidos.append({**row, "periodo": periodo, "motivo": "Factura ya cobrada"})
			continue

		mode = _resolve_mode_of_payment(row["notas"], row["cobrador"])
		entry = {
			**row,
			"periodo": periodo,
			"sales_invoice": invoice_name,
			"mode_of_payment": mode,
			"outstanding": outstanding,
		}

		if dry_run:
			ok.append(entry)
			continue

		try:
			pe_name = registrar_cobro_manual(
				socio_name,
				invoice_name,
				mode_of_payment=mode,
				posting_date=row["fecha"],
			)
			ok.append({**entry, "payment_entry": pe_name})
		except Exception as exc:
			errores.append({**entry, "error": str(exc)})

	result = {
		"operacion": "import_cobranza_excel",
		"ejecutado_en": datetime.now().isoformat(timespec="seconds"),
		"site": frappe.local.site,
		"dry_run": dry_run,
		"file_path": file_path,
		"fecha_desde": str(desde) if desde else None,
		"fecha_hasta": str(hasta) if hasta else None,
		"resumen": {
			"procesados": len(ok),
			"omitidos": len(omitidos),
			"errores": len(errores),
		},
		"procesados": len(ok),
		"omitidos": len(omitidos),
		"errores": len(errores),
		"ok": [_serialize_row(r) for r in ok],
		"omitidos_detalle": [_serialize_row(r) for r in omitidos],
		"errores_detalle": [_serialize_row(r) for r in errores],
	}

	if not log_path:
		stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
		mode = "dry-run" if dry_run else "apply"
		log_path = str(
			Path(frappe.get_site_path("private", "logs", "cobranza"))
			/ f"import-cobranza-{mode}-{stamp}.json"
		)

	result["log_path"] = log_path
	written = _write_log(log_path, result)
	frappe.logger("club_management.cobranza").info("import_cobranza_excel %s", result["resumen"])
	frappe.msgprint(frappe.as_json({"resumen": result["resumen"], "log_path": written}, indent=2))
	return result
