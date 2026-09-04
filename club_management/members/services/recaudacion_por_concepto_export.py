"""Export Excel/PDF de rendición de cobranza (alineado al modelo Secretaría).

Spec: `club_management/specs/informe_rendicion_cobranza_secretaria.md` (fase 4)
"""

from __future__ import annotations

from datetime import date
from io import BytesIO
from typing import Any

import frappe
from frappe import _
from frappe.utils import flt, format_datetime, formatdate, getdate, now_datetime, today

from club_management.members.services.concepto_informe_label import agrupacion_tipo_concepto
from club_management.members.services.modos_pago_desk import DESK_MODOS_PAGO_COBRANZA
from club_management.members.services.recibo_pago import get_recibo_config
from club_management.members.services.recaudacion_por_concepto import (
	VISTA_PAGOS_DIA,
	get_informe_recaudacion_por_concepto,
)

SHEET_RESUMEN = "Resumen Ejecutivo CD"
SHEET_DETALLE = "Detalle_Cobranzas"


def _label_medio(mode: str) -> str:
	for row in DESK_MODOS_PAGO_COBRANZA:
		if row["value"] == mode:
			return _(row["label"])
	return mode


def filters_for_export(filters: dict[str, Any] | str | None) -> dict[str, Any]:
	"""Normaliza filtros Desk (incluye vista Pagos del día → un solo día)."""
	raw: dict[str, Any]
	if isinstance(filters, str):
		raw = frappe.parse_json(filters) or {}
	else:
		raw = dict(filters or {})

	if (raw.get("vista") or "") == VISTA_PAGOS_DIA:
		fecha = raw.get("fecha") or raw.get("fecha_desde") or today()
		return {
			"fecha_desde": fecha,
			"fecha_hasta": fecha,
			"periodo_cobro": (raw.get("periodo_cobro") or "").strip(),
			"agrupacion": (raw.get("agrupacion") or "").strip(),
			"solo_cuotas_sociales": bool(raw.get("solo_cuotas_sociales")),
			"medio_pago": (raw.get("medio_pago") or "").strip(),
		}
	return {
		"fecha_desde": raw.get("fecha_desde") or today(),
		"fecha_hasta": raw.get("fecha_hasta") or today(),
		"periodo_cobro": (raw.get("periodo_cobro") or "").strip(),
		"agrupacion": (raw.get("agrupacion") or "").strip(),
		"solo_cuotas_sociales": bool(raw.get("solo_cuotas_sociales")),
		"medio_pago": (raw.get("medio_pago") or "").strip(),
	}


def _pct(parte: float, total: float) -> float:
	if flt(total) <= 0:
		return 0.0
	return round(100.0 * flt(parte) / flt(total), 6)


def _club_nombre() -> str:
	try:
		return str(get_recibo_config()["encabezado"]["institucion_nombre"])
	except Exception:
		return "Institucion Cultural y Deportiva Pedro Echague"


def build_rendicion_export_context(informe: dict[str, Any]) -> dict[str, Any]:
	"""Enriquece el informe con KPIs, conteos y % para Excel/PDF."""
	lineas = list(informe.get("lineas") or [])
	total = flt(informe.get("total"))
	ops_concepto: dict[str, int] = {}
	ops_medio: dict[str, int] = {}
	for row in lineas:
		concepto = str(row.get("concepto_informe") or _("Sin concepto"))
		mode = str(row.get("mode_of_payment") or _("Sin medio"))
		ops_concepto[concepto] = ops_concepto.get(concepto, 0) + 1
		ops_medio[mode] = ops_medio.get(mode, 0) + 1

	por_concepto: list[dict[str, Any]] = []
	total_cuotas = 0.0
	for row in informe.get("por_concepto") or []:
		concepto = str(row["concepto"])
		monto = flt(row["total"])
		if agrupacion_tipo_concepto(concepto) == "Cuota":
			total_cuotas += monto
		por_concepto.append(
			{
				"concepto": concepto,
				"total": monto,
				"operaciones": ops_concepto.get(concepto, 0),
				"pct": _pct(monto, total),
			}
		)

	por_medio: list[dict[str, Any]] = []
	for row in informe.get("por_medio") or []:
		mode = str(row["mode_of_payment"])
		monto = flt(row["total"])
		por_medio.append(
			{
				"mode_of_payment": mode,
				"label": _label_medio(mode),
				"total": monto,
				"operaciones": ops_medio.get(mode, 0),
				"pct": _pct(monto, total),
			}
		)

	club_nombre = _club_nombre()
	fecha_desde = getdate(informe.get("fecha_desde") or today())
	fecha_hasta = getdate(informe.get("fecha_hasta") or today())
	if fecha_desde == fecha_hasta:
		rango_txt = _("Fecha de Cobro: {0}").format(formatdate(fecha_desde))
	else:
		rango_txt = _("Período de Cobro: {0} — {1}").format(
			formatdate(fecha_desde),
			formatdate(fecha_hasta),
		)

	detalle_rows: list[dict[str, Any]] = []
	for idx, row in enumerate(lineas, start=1):
		detalle_rows.append(
			{
				"nro": idx,
				"concepto": row.get("concepto_informe") or "",
				"fecha_cobro": row.get("posting_date"),
				"nro_socio": row.get("nro_socio") or "",
				"socio_label": row.get("socio_label") or "",
				"periodo": row.get("periodo") or "",
				"medio": _label_medio(str(row.get("mode_of_payment") or "")),
				"monto": flt(row.get("paid_amount")),
				"comprobante": row.get("payment_entry") or "",
			}
		)

	return {
		"club_nombre": club_nombre,
		"titulo": _("{0} — REPORTE DE RECAUDACIÓN Y COBRANZAS").format(club_nombre.upper()),
		"subtitulo": _("Rendición de ingresos por caja y canales digitales | {0}").format(rango_txt),
		"generado": now_datetime(),
		"fecha_desde": fecha_desde,
		"fecha_hasta": fecha_hasta,
		"total": total,
		"operaciones": len(lineas),
		"total_cuotas": flt(total_cuotas),
		"total_aranceles": flt(total - total_cuotas),
		"por_concepto": por_concepto,
		"por_medio": por_medio,
		"detalle": detalle_rows,
	}


def _filename_stem(fecha_desde: date, fecha_hasta: date) -> str:
	slug = "Cobranzas"
	if fecha_desde == fecha_hasta:
		return f"Reporte_{slug}_{fecha_desde.isoformat()}"
	return f"Reporte_{slug}_{fecha_desde.isoformat()}_{fecha_hasta.isoformat()}"


def build_rendicion_xlsx_bytes(informe: dict[str, Any]) -> bytes:
	"""Genera el `.xlsx` de 2 hojas según el modelo de rendición."""
	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Font

	ctx = build_rendicion_export_context(informe)
	wb = Workbook()
	ws = wb.active
	ws.title = SHEET_RESUMEN

	bold = Font(bold=True, size=12)
	title_font = Font(bold=True, size=14)
	header_font = Font(bold=True, size=10)
	money_format = "#,##0.00"
	pct_format = "0.00%"

	ws["B2"] = ctx["titulo"]
	ws["B2"].font = title_font
	ws["B3"] = ctx["subtitulo"]
	ws["B4"] = _("Destinatario: Comisión Directiva / Tesorería | Generado: {0}").format(
		format_datetime(ctx["generado"])
	)

	ws["B6"] = _("RECAUDACIÓN TOTAL")
	ws["D6"] = _("OPERACIONES")
	ws["F6"] = _("TOTAL CUOTAS SOCIALES")
	ws["H6"] = _("TOTAL ARANCELES / DEPORTES")
	for cell in ("B6", "D6", "F6", "H6"):
		ws[cell].font = header_font

	ws["B7"] = ctx["total"]
	ws["B7"].number_format = money_format
	ws["B7"].font = bold
	ws["D7"] = ctx["operaciones"]
	ws["D7"].font = bold
	ws["F7"] = ctx["total_cuotas"]
	ws["F7"].number_format = money_format
	ws["F7"].font = bold
	ws["H7"] = ctx["total_aranceles"]
	ws["H7"].number_format = money_format
	ws["H7"].font = bold

	ws["B10"] = _("1. RESUMEN POR CONCEPTO Y RAMA DEPORTIVA")
	ws["B10"].font = bold
	ws["G10"] = _("2. ARQUEO POR MEDIO DE COBRO")
	ws["G10"].font = bold

	ws["B11"] = _("Concepto / Cuenta")
	ws["C11"] = _("Cant. Operaciones")
	ws["D11"] = _("Total Recaudado ($)")
	ws["E11"] = _("% Participación")
	ws["G11"] = _("Medio de Cobro")
	ws["H11"] = _("Operaciones")
	ws["I11"] = _("Total Recaudado ($)")
	ws["J11"] = _("% Total")
	for col in ("B", "C", "D", "E", "G", "H", "I", "J"):
		ws[f"{col}11"].font = header_font

	row = 12
	for item in ctx["por_concepto"]:
		ws.cell(row, 2, item["concepto"])
		ws.cell(row, 3, item["operaciones"])
		cell_total = ws.cell(row, 4, item["total"])
		cell_total.number_format = money_format
		cell_pct = ws.cell(row, 5, flt(item["pct"]) / 100.0)
		cell_pct.number_format = pct_format
		row += 1

	total_row = row
	ws.cell(total_row, 2, _("TOTAL GENERAL CONCEPTOS")).font = bold
	ws.cell(total_row, 3, ctx["operaciones"]).font = bold
	cell_tg = ws.cell(total_row, 4, ctx["total"])
	cell_tg.number_format = money_format
	cell_tg.font = bold
	cell_tp = ws.cell(total_row, 5, 1.0 if ctx["total"] else 0.0)
	cell_tp.number_format = pct_format
	cell_tp.font = bold

	row_m = 12
	for item in ctx["por_medio"]:
		ws.cell(row_m, 7, item["label"])
		ws.cell(row_m, 8, item["operaciones"])
		cell_m = ws.cell(row_m, 9, item["total"])
		cell_m.number_format = money_format
		cell_mp = ws.cell(row_m, 10, flt(item["pct"]) / 100.0)
		cell_mp.number_format = pct_format
		row_m += 1

	ws.cell(row_m, 7, _("TOTAL POR MEDIOS")).font = bold
	ws.cell(row_m, 8, ctx["operaciones"]).font = bold
	cell_tm = ws.cell(row_m, 9, ctx["total"])
	cell_tm.number_format = money_format
	cell_tm.font = bold
	cell_tmp = ws.cell(row_m, 10, 1.0 if ctx["total"] else 0.0)
	cell_tmp.number_format = pct_format
	cell_tmp.font = bold

	firmas_row = max(total_row, row_m) + 3
	ws.cell(firmas_row, 7, _("3. CONCILIACIÓN Y FIRMAS DE CONTROL")).font = bold
	ws.cell(firmas_row + 2, 7, _("Efectivo Rendido en Tesorería:"))
	efectivo = next(
		(item["total"] for item in ctx["por_medio"] if item["mode_of_payment"] == "Cash"),
		0.0,
	)
	cell_ef = ws.cell(firmas_row + 2, 9, efectivo)
	cell_ef.number_format = money_format
	ws.cell(firmas_row + 3, 7, _("Firma Responsable Caja / Administración:"))
	ws.cell(firmas_row + 3, 9, "_______________________")
	ws.cell(firmas_row + 4, 7, _("Firma Tesorería / Comisión Directiva:"))
	ws.cell(firmas_row + 4, 9, "_______________________")

	for col, width in (
		("B", 36),
		("C", 16),
		("D", 18),
		("E", 14),
		("G", 28),
		("H", 14),
		("I", 18),
		("J", 12),
	):
		ws.column_dimensions[col].width = width

	# Detalle
	wd = wb.create_sheet(SHEET_DETALLE)
	headers = [
		"#",
		_("Concepto"),
		_("Fecha Cobro"),
		_("N° Socio"),
		_("Apellido y Nombre Socio"),
		_("Período Imputado"),
		_("Medio de Pago"),
		_("Monto Cobrado ($)"),
		_("N° Comprobante / Recibo"),
	]
	for col, header in enumerate(headers, start=1):
		cell = wd.cell(1, col, header)
		cell.font = header_font
		cell.alignment = Alignment(wrap_text=True)

	for idx, row in enumerate(ctx["detalle"], start=1):
		r = idx + 1
		wd.cell(r, 1, row["nro"])
		wd.cell(r, 2, row["concepto"])
		fecha = row["fecha_cobro"]
		wd.cell(r, 3, formatdate(fecha) if fecha else "")
		wd.cell(r, 4, row["nro_socio"])
		wd.cell(r, 5, row["socio_label"])
		wd.cell(r, 6, row["periodo"])
		wd.cell(r, 7, row["medio"])
		cell_monto = wd.cell(r, 8, row["monto"])
		cell_monto.number_format = money_format
		wd.cell(r, 9, row["comprobante"])

	total_r = len(ctx["detalle"]) + 2
	wd.cell(total_r, 1, "TOTAL").font = bold
	cell_tot = wd.cell(total_r, 8, ctx["total"])
	cell_tot.number_format = money_format
	cell_tot.font = bold
	wd.cell(total_r, 9, _("{0} cobranzas").format(ctx["operaciones"]))

	from openpyxl.utils import get_column_letter

	for col, width in enumerate((6, 32, 12, 12, 36, 14, 18, 16, 22), start=1):
		wd.column_dimensions[get_column_letter(col)].width = width

	buf = BytesIO()
	wb.save(buf)
	return buf.getvalue()


def build_rendicion_pdf_html(informe: dict[str, Any]) -> str:
	ctx = build_rendicion_export_context(informe)
	fmt_money = frappe.format_value
	por_concepto = [
		{
			**row,
			"total_txt": fmt_money(row["total"], {"fieldtype": "Currency"}),
			"pct_txt": f"{flt(row['pct']):.2f}%",
		}
		for row in ctx["por_concepto"]
	]
	por_medio = [
		{
			**row,
			"total_txt": fmt_money(row["total"], {"fieldtype": "Currency"}),
			"pct_txt": f"{flt(row['pct']):.2f}%",
		}
		for row in ctx["por_medio"]
	]
	detalle = [
		{
			**row,
			"fecha_txt": formatdate(row["fecha_cobro"]) if row.get("fecha_cobro") else "",
			"monto_txt": fmt_money(row["monto"], {"fieldtype": "Currency"}),
		}
		for row in ctx["detalle"]
	]
	return frappe.render_template(
		"templates/recaudacion_por_concepto_pdf.html",
		{
			"titulo": ctx["titulo"],
			"subtitulo": ctx["subtitulo"],
			"generado_txt": format_datetime(ctx["generado"]),
			"fecha_desde_txt": formatdate(ctx["fecha_desde"]),
			"fecha_hasta_txt": formatdate(ctx["fecha_hasta"]),
			"total_txt": fmt_money(ctx["total"], {"fieldtype": "Currency"}),
			"operaciones": ctx["operaciones"],
			"total_cuotas_txt": fmt_money(ctx["total_cuotas"], {"fieldtype": "Currency"}),
			"total_aranceles_txt": fmt_money(ctx["total_aranceles"], {"fieldtype": "Currency"}),
			"por_concepto": por_concepto,
			"por_medio": por_medio,
			"detalle": detalle,
		},
	)


def build_rendicion_pdf_bytes(informe: dict[str, Any]) -> bytes:
	from frappe.utils.pdf import get_pdf

	html = build_rendicion_pdf_html(informe)
	return get_pdf(html)


def export_filename(informe: dict[str, Any], extension: str) -> str:
	fecha_desde = getdate(informe.get("fecha_desde") or today())
	fecha_hasta = getdate(informe.get("fecha_hasta") or today())
	return f"{_filename_stem(fecha_desde, fecha_hasta)}.{extension}"


def build_export_bytes(
	filters: dict[str, Any] | str | None,
	*,
	file_format: str,
) -> tuple[str, bytes]:
	"""Devuelve (filename, content) para Excel o PDF."""
	parsed = filters_for_export(filters)
	informe = get_informe_recaudacion_por_concepto(parsed)
	fmt = (file_format or "Excel").strip().lower()
	if fmt in {"pdf"}:
		return export_filename(informe, "pdf"), build_rendicion_pdf_bytes(informe)
	return export_filename(informe, "xlsx"), build_rendicion_xlsx_bytes(informe)
