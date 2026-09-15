"""Tests: export Excel/PDF de Recaudación por concepto.

Spec: `club_management/specs/informe_rendicion_cobranza_secretaria.md` (fase 4)
"""

from __future__ import annotations

from io import BytesIO
from unittest import TestCase

import frappe
from frappe.utils import flt

from club_management.integrations.payment_ledger_postgres import apply_patch
from club_management.members.services.cobranza_manual import (
	SALES_INVOICE_DOCTYPE,
	_campo_periodo_cobro,
	_campo_socio_en,
	erpnext_cobranza_disponible,
	registrar_cobro_manual,
	sync_saldo_deuda_socio,
)
from club_management.members.services.cuotas_sociales_setup import sync_cuotas_sociales_club
from club_management.members.services.recaudacion_por_concepto import (
	get_informe_recaudacion_por_concepto,
)
from club_management.members.services.recaudacion_por_concepto_export import (
	SHEET_DETALLE,
	SHEET_RESUMEN,
	build_rendicion_export_context,
	build_rendicion_pdf_html,
	build_rendicion_xlsx_bytes,
	filters_for_export,
)
from club_management.members.services.socio_transitions import cambiar_estado
from club_management.members.test_helpers import MembersTestCase, insert_socio, make_secretaria_user


def _sample_informe() -> dict:
	return {
		"fecha_desde": "2026-09-01",
		"fecha_hasta": "2026-09-01",
		"lineas": [
			{
				"concepto_informe": "Cuota Social Menor",
				"posting_date": "2026-09-01",
				"nro_socio": "10975",
				"socio_label": "GILARDI, IGNACIO TOMAS",
				"apellido": "GILARDI",
				"nombre": "IGNACIO TOMAS",
				"periodo": "09/2026",
				"mode_of_payment": "Cash",
				"paid_amount": 28500.0,
				"payment_entry": "ACC-PAY-2026-01984",
				"sales_invoice": "ACC-SINV-0001",
			},
			{
				"concepto_informe": "CTO COMP BASQ TIRA A/B/FLEX",
				"posting_date": "2026-09-01",
				"nro_socio": "10975",
				"socio_label": "GILARDI, IGNACIO TOMAS",
				"apellido": "GILARDI",
				"nombre": "IGNACIO TOMAS",
				"periodo": "09/2026",
				"mode_of_payment": "Wire Transfer",
				"paid_amount": 6000.0,
				"payment_entry": "ACC-PAY-2026-01984",
				"sales_invoice": "ACC-SINV-0002",
			},
			{
				"concepto_informe": "Cuota Social Menor",
				"posting_date": "2026-09-01",
				"nro_socio": "12246",
				"socio_label": "ZAFFORE, GUADALUPE",
				"apellido": "ZAFFORE",
				"nombre": "GUADALUPE",
				"periodo": "09/2026",
				"mode_of_payment": "Wire Transfer",
				"paid_amount": 27500.0,
				"payment_entry": "ACC-PAY-2026-01978",
				"sales_invoice": "ACC-SINV-0003",
			},
		],
		"por_concepto": [
			{"concepto": "Cuota Social Menor", "total": 56000.0},
			{"concepto": "CTO COMP BASQ TIRA A/B/FLEX", "total": 6000.0},
		],
		"por_medio": [
			{"mode_of_payment": "Cash", "total": 28500.0},
			{"mode_of_payment": "Wire Transfer", "total": 33500.0},
		],
		"total": 62000.0,
	}


class TestRecaudacionExportBuilders(TestCase):
	def test_context_kpis_y_porcentajes(self) -> None:
		ctx = build_rendicion_export_context(_sample_informe())
		self.assertEqual(ctx["operaciones"], 3)
		self.assertAlmostEqual(flt(ctx["total"]), 62000.0)
		self.assertAlmostEqual(flt(ctx["total_cuotas"]), 56000.0)
		self.assertAlmostEqual(flt(ctx["total_aranceles"]), 6000.0)

		por_concepto = {row["concepto"]: row for row in ctx["por_concepto"]}
		self.assertEqual(por_concepto["Cuota Social Menor"]["operaciones"], 2)
		self.assertAlmostEqual(flt(por_concepto["Cuota Social Menor"]["pct"]), 90.322580645, places=4)

		por_medio = {row["mode_of_payment"]: row for row in ctx["por_medio"]}
		self.assertIn("Efectivo", por_medio["Cash"]["label"])
		self.assertEqual(por_medio["Cash"]["operaciones"], 1)

	def test_filters_pagos_dia_convierten_a_rango(self) -> None:
		parsed = filters_for_export({"vista": "Pagos del día", "fecha": "2026-09-01"})
		self.assertEqual(str(parsed["fecha_desde"]), "2026-09-01")
		self.assertEqual(str(parsed["fecha_hasta"]), "2026-09-01")

	def test_xlsx_dos_hojas_y_detalle(self) -> None:
		import openpyxl

		content = build_rendicion_xlsx_bytes(_sample_informe())
		self.assertTrue(content[:2] == b"PK")
		wb = openpyxl.load_workbook(BytesIO(content))
		self.assertEqual(wb.sheetnames, [SHEET_RESUMEN, SHEET_DETALLE])

		detalle = wb[SHEET_DETALLE]
		headers = [detalle.cell(1, c).value for c in range(1, 10)]
		self.assertEqual(
			headers,
			[
				"#",
				"Concepto",
				"Fecha Cobro",
				"N° Socio",
				"Apellido y Nombre Socio",
				"Período Imputado",
				"Medio de Pago",
				"Monto Cobrado ($)",
				"N° Comprobante / Recibo",
			],
		)
		self.assertEqual(detalle.cell(2, 2).value, "Cuota Social Menor")
		self.assertEqual(detalle.cell(2, 9).value, "ACC-PAY-2026-01984")
		self.assertEqual(detalle.cell(5, 1).value, "TOTAL")
		self.assertAlmostEqual(flt(detalle.cell(5, 8).value), 62000.0)

		resumen = wb[SHEET_RESUMEN]
		joined = " ".join(
			str(resumen.cell(r, c).value or "")
			for r in range(1, 30)
			for c in range(1, 12)
		)
		self.assertIn("REPORTE DE RECAUDACIÓN", joined)
		self.assertIn("Cuota Social Menor", joined)
		self.assertIn("ARQUEO POR MEDIO", joined)

	def test_pdf_html_incluye_resumen_y_detalle(self) -> None:
		html = build_rendicion_pdf_html(_sample_informe())
		self.assertIn("Recaudación total", html)
		self.assertIn("Cuota Social Menor", html)
		self.assertIn("ACC-PAY-2026-01984", html)
		self.assertIn("Detalle", html)


class TestRecaudacionExportApi(MembersTestCase):
	_DESDE = "2020-05-01"
	_HASTA = "2020-05-31"
	_DIA = "2020-05-10"
	_PERIODO = "05/2020"

	def setUp(self) -> None:
		super().setUp()
		if not erpnext_cobranza_disponible():
			self.skipTest("ERPNext Sales Invoice no instalado")
		apply_patch()
		self._secretaria = make_secretaria_user("secretaria.rec.export@example.com")
		sync_cuotas_sociales_club()
		settings = frappe.get_single("Club Settings")
		if not settings.company:
			company = frappe.db.get_value("Company", {}, "name")
			if company:
				settings.company = company
				settings.save(ignore_permissions=True)

	def _ensure_item(self, code: str, rate: float) -> str:
		if not frappe.db.exists("Item", code):
			frappe.get_doc(
				{
					"doctype": "Item",
					"item_code": code,
					"item_name": code,
					"item_group": frappe.db.get_value("Item Group", {"is_group": 0}, "name")
					or "Products",
					"stock_uom": "Nos",
					"is_sales_item": 1,
					"is_stock_item": 0,
					"standard_rate": rate,
				}
			).insert(ignore_permissions=True)
		return code

	def _crear_si(self, socio_name: str, item_code: str, rate: float, description: str) -> str:
		from club_management.members.services.cobranza_manual import (
			_default_company,
			ensure_customer_for_socio,
		)

		campo = _campo_socio_en(SALES_INVOICE_DOCTYPE)
		campo_periodo = _campo_periodo_cobro()
		payload: dict = {
			"doctype": SALES_INVOICE_DOCTYPE,
			"customer": ensure_customer_for_socio(socio_name),
			"company": _default_company(),
			"posting_date": self._DIA,
			"due_date": self._DIA,
			"set_posting_time": 1,
			campo: socio_name,
			"items": [
				{"item_code": item_code, "qty": 1, "rate": rate, "description": description}
			],
		}
		if campo_periodo:
			payload[campo_periodo] = self._PERIODO
		doc = frappe.get_doc(payload)
		doc.insert(ignore_permissions=True)
		doc.submit()
		return doc.name

	def test_export_excel_api_totales_y_permisos(self) -> None:
		from club_management.members.api.cobranza_desk import export_recaudacion_por_concepto

		socio = insert_socio(dni="99450001", email="rec.export@example.com", categoria="Activo")
		cambiar_estado(socio.name, "Activo", motivo="Test export recaudación")
		item = self._ensure_item("ICDPE-CUOTA-EXPORT-TEST", 15000)
		si = self._crear_si(socio.name, item, 15000, "Cuota Social Activo")
		sync_saldo_deuda_socio(socio.name)
		esperado = flt(frappe.db.get_value(SALES_INVOICE_DOCTYPE, si, "outstanding_amount"))

		frappe.set_user(self._secretaria)
		try:
			registrar_cobro_manual(
				socio.name,
				si,
				mode_of_payment="Cash",
				posting_date=self._DIA,
			)
			filters = {"fecha_desde": self._DESDE, "fecha_hasta": self._HASTA}
			informe = get_informe_recaudacion_por_concepto(filters)
			export_recaudacion_por_concepto(filters=filters, file_format="Excel")
		finally:
			frappe.set_user("Administrator")

		self.assertAlmostEqual(flt(informe["total"]), esperado)
		self.assertEqual(frappe.response.get("type"), "binary")
		self.assertTrue(str(frappe.response.get("filename") or "").endswith(".xlsx"))
		self.assertTrue((frappe.response.get("filecontent") or b"")[:2] == b"PK")

		guest = frappe.get_doc(
			{
				"doctype": "User",
				"email": "rec.export.guest@example.com",
				"first_name": "GuestExport",
				"send_welcome_email": 0,
			}
		)
		if not frappe.db.exists("User", guest.email):
			guest.insert(ignore_permissions=True)
			guest.add_roles("Employee")

		frappe.set_user(guest.email)
		try:
			with self.assertRaises(frappe.PermissionError):
				export_recaudacion_por_concepto(
					filters={"fecha_desde": self._DESDE, "fecha_hasta": self._HASTA},
					file_format="Excel",
				)
		finally:
			frappe.set_user("Administrator")
