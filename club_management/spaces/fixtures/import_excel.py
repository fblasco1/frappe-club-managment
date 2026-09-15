"""Importación de fixtures / partidos desde Excel de ligas."""

from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Any

import frappe
from frappe import _
from frappe.utils.file_manager import get_file_path

from club_management.spaces.fixtures.contract import ORIGIN_LIGA_EXCEL, LOCALIA_LOCAL
from club_management.spaces.fixtures.espacio_map import resolve_espacio
from club_management.spaces.fixtures.parse import normalize_partido, parse_hora_hasta_explicita
from club_management.spaces.fixtures.upsert import import_fixture_rows

try:
	import openpyxl
except Exception:  # pragma: no cover
	openpyxl = None  # type: ignore[assignment]

REQUIRED_HEADERS = frozenset(
	{"fecha", "hora", "hora_hasta", "espacio", "categoria", "tira", "rival"}
)
HEADER_ALIASES: dict[str, str] = {
	"fecha": "fecha",
	"hora": "hora",
	"hora_inicio": "hora",
	"hora_desde": "hora",
	"hora_fin": "hora_hasta",
	"hora_hasta": "hora_hasta",
	"espacio": "espacio",
	"categoria": "categoria",
	"categoría": "categoria",
	"tira": "tira",
	"equipo": "equipo",
	"rival": "rival",
	"localia": "localia",
	"localía": "localia",
	"id_externo": "id_externo",
	"external_id": "id_externo",
	"id_partido": "id_externo",
	"origen": "origen",
	"source": "origen",
	"motivo": "motivo",
}


def _normalize_header(value: Any) -> str:
	text = str(value or "").strip().lower()
	return HEADER_ALIASES.get(text, text)


def _cell_fecha(value: Any) -> str:
	from datetime import date, datetime

	if value is None:
		return ""
	if isinstance(value, datetime):
		return value.date().isoformat()
	if isinstance(value, date):
		return value.isoformat()
	text = str(value).strip()
	if " " in text and re.match(r"^\d{4}-\d{2}-\d{2}\s", text):
		return text.split(" ", 1)[0]
	return text


def _cell_hora(value: Any) -> str:
	from datetime import datetime, time

	if value is None:
		return ""
	if isinstance(value, datetime):
		return value.strftime("%H:%M")
	if isinstance(value, time):
		return value.strftime("%H:%M")
	return str(value).strip()


def _cell_str(value: Any) -> str:
	if value is None:
		return ""
	return str(value).strip()


def _identity_value(value: Any) -> str:
	return " ".join(str(value or "").strip().casefold().split())


def build_excel_external_id(raw: dict[str, Any]) -> str:
	"""Genera clave idempotente sin exponer IDs técnicos en la planilla."""
	identity = "|".join(
		_identity_value(raw.get(field))
		for field in ("fecha", "categoria", "tira", "rival")
	)
	return f"xlsx-{hashlib.sha256(identity.encode('utf-8')).hexdigest()[:20]}"


def build_excel_payload(raw: dict[str, Any]) -> dict[str, Any] | str:
	required_values = {
		"fecha": "fecha",
		"hora": "hora_inicio",
		"hora_hasta": "hora_fin",
		"espacio": "espacio",
		"categoria": "categoria",
		"tira": "tira",
		"rival": "rival",
	}
	missing = [label for field, label in required_values.items() if not raw.get(field)]
	if missing:
		return _("Faltan valores obligatorios: {0}").format(", ".join(missing))
	if not parse_hora_hasta_explicita(str(raw.get("hora_hasta") or "")):
		return _("Hora de fin inválida: {0}").format(repr(raw.get("hora_hasta")))

	categoria = str(raw.get("categoria") or "").strip()
	tira = str(raw.get("tira") or "").strip()
	return {
		"source": ORIGIN_LIGA_EXCEL,
		"external_id": build_excel_external_id(raw),
		"fecha": raw.get("fecha") or "",
		"hora": raw.get("hora") or "",
		"hora_hasta": raw.get("hora_hasta") or "",
		"espacio": raw.get("espacio") or None,
		"categoria": categoria,
		"tira": tira,
		"equipo": " ".join(value for value in (tira, categoria) if value),
		"rival": raw.get("rival") or "",
		"localia": "Local",
		"motivo": raw.get("motivo") or "",
	}


def resolve_excel_path(file_url: str | Path) -> Path:
	"""Resuelve path local de un File Desk o ruta absoluta."""
	raw = str(file_url or "").strip()
	if not raw:
		frappe.throw(_("Falta archivo Excel"), frappe.ValidationError)
	path = Path(raw)
	if path.is_file():
		resolved_path = path
	else:
		try:
			resolved_path = Path(get_file_path(raw))
		except Exception as exc:
			frappe.throw(_("No se pudo leer el archivo: {0}").format(exc), frappe.ValidationError)
	if resolved_path.suffix.lower() != ".xlsx":
		frappe.throw(_("Solo se admiten archivos .xlsx"), frappe.ValidationError)
	if not resolved_path.is_file():
		frappe.throw(_("Archivo no encontrado: {0}").format(raw), frappe.ValidationError)
	return resolved_path


def read_excel_rows(path: str | Path) -> tuple[list[dict[str, Any]], list[str]]:
	"""Lee filas de la hoja Fixture (o la primera). Devuelve (rows, errores_globales)."""
	if openpyxl is None:
		frappe.throw(
			_("Dependencia openpyxl no disponible para importar Excel"),
			frappe.ValidationError,
		)
	wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
	sheet = wb["Fixture"] if "Fixture" in wb.sheetnames else wb[wb.sheetnames[0]]
	rows_iter = sheet.iter_rows(values_only=True)
	try:
		header_raw = next(rows_iter)
	except StopIteration:
		return [], [_("Excel vacío")]
	headers = [_normalize_header(h) for h in header_raw]
	missing = REQUIRED_HEADERS - {h for h in headers if h}
	if missing:
		return [], [_("Faltan columnas obligatorias: {0}").format(", ".join(sorted(missing)))]

	rows: list[dict[str, Any]] = []
	for idx, values in enumerate(rows_iter, start=2):
		if values is None or all(v is None or str(v).strip() == "" for v in values):
			continue
		raw: dict[str, Any] = {}
		for col_idx, header in enumerate(headers):
			if not header or header.startswith("_"):
				continue
			cell = values[col_idx] if col_idx < len(values) else None
			if header == "fecha":
				raw[header] = _cell_fecha(cell)
			elif header in {"hora", "hora_hasta"}:
				raw[header] = _cell_hora(cell)
			else:
				raw[header] = _cell_str(cell)
		raw["_row"] = idx
		rows.append(raw)
	return rows, []


def preview_excel_fixtures(file_url: str | Path) -> dict[str, Any]:
	"""Valida filas sin escribir BD."""
	path = resolve_excel_path(file_url)
	raw_rows, global_errors = read_excel_rows(path)
	filas_ok: list[dict[str, Any]] = []
	errores: list[str] = list(global_errors)

	for raw in raw_rows:
		row_no = raw.get("_row")
		payload = build_excel_payload(raw)
		if isinstance(payload, str):
			errores.append(f"Fila {row_no}: {payload}")
			continue
		normalized = normalize_partido(payload)
		if isinstance(normalized, str):
			errores.append(f"Fila {row_no}: {normalized}")
			continue
		if (
			normalized.localia == LOCALIA_LOCAL
			and normalized.source == ORIGIN_LIGA_EXCEL
			and not (normalized.espacio or "").strip()
		):
			errores.append(f"Fila {row_no}: falta espacio (id {normalized.external_id})")
			continue
		espacio = resolve_espacio(normalized)
		if normalized.localia == LOCALIA_LOCAL and not espacio:
			errores.append(f"Fila {row_no}: sin espacio mapeado (id {normalized.external_id})")
			continue
		filas_ok.append(
			{
				"row": row_no,
				"external_id": normalized.external_id,
				"fecha": normalized.fecha,
				"hora_desde": normalized.hora_desde,
				"hora_hasta": normalized.hora_hasta,
				"espacio": espacio or normalized.espacio,
				"motivo": normalized.motivo,
				"localia": normalized.localia,
				"origen": normalized.source,
			}
		)

	return {
		"filas_ok": filas_ok,
		"errores": errores,
		"total_filas": len(raw_rows),
		"path": str(path),
	}


def apply_excel_fixtures(
	file_url: str | Path,
	*,
	cancel_missing: bool = False,
) -> dict[str, Any]:
	"""Importa Excel → upsert Reserva Espacio."""
	path = resolve_excel_path(file_url)
	raw_rows, global_errors = read_excel_rows(path)
	payload_rows: list[dict[str, Any]] = []
	omitidos: list[str] = list(global_errors)

	for raw in raw_rows:
		row_no = raw.get("_row")
		payload = build_excel_payload(raw)
		if isinstance(payload, str):
			omitidos.append(f"Fila {row_no}: {payload}")
			continue
		# Validación previa para mensajes con número de fila
		normalized = normalize_partido(payload)
		if isinstance(normalized, str):
			omitidos.append(f"Fila {row_no}: {normalized}")
			continue
		if (
			normalized.localia == LOCALIA_LOCAL
			and normalized.source == ORIGIN_LIGA_EXCEL
			and not (normalized.espacio or "").strip()
		):
			omitidos.append(f"Fila {row_no}: falta espacio (id {normalized.external_id})")
			continue
		if normalized.localia == LOCALIA_LOCAL and not resolve_espacio(normalized):
			omitidos.append(f"Fila {row_no}: sin espacio mapeado (id {normalized.external_id})")
			continue
		payload_rows.append(payload)

	result = import_fixture_rows(
		payload_rows,
		source=None,
		cancel_missing=cancel_missing,
	)
	result["omitidos"] = omitidos + list(result.get("omitidos") or [])
	result["path"] = str(path)
	return result
