"""Tests importación fixtures / partidos desde Excel de ligas."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import frappe
import openpyxl
from frappe.utils.file_manager import save_file

from club_management.members.test_helpers import MembersTestCase
from club_management.spaces.fixtures.contract import ORIGIN_LIGA_EXCEL
from club_management.spaces.fixtures.import_excel import (
	apply_excel_fixtures,
	build_excel_external_id,
	preview_excel_fixtures,
)
from club_management.spaces.fixtures.upsert import find_reserva_by_fixture
from club_management.spaces.helpers import insert_espacio, make_coordinacion_user
from club_management.spaces.import_horarios import CANCHA_2, CANCHA_3


def _write_sample_xlsx(path: Path, rows: list[list[object]]) -> Path:
	wb = openpyxl.Workbook()
	ws = wb.active
	ws.title = "Fixture"
	for row in rows:
		ws.append(row)
	path.parent.mkdir(parents=True, exist_ok=True)
	wb.save(path)
	return path


_HEADERS = [
	"fecha",
	"hora_inicio",
	"hora_fin",
	"espacio",
	"categoria",
	"tira",
	"rival",
]


class TestFixtureExcel(MembersTestCase):
	def setUp(self) -> None:
		super().setUp()
		frappe.set_user("Administrator")
		for name in (CANCHA_2, CANCHA_3):
			if not frappe.db.exists("Espacio", name):
				insert_espacio(name, tipo="Cancha")
		self._tmp = Path(frappe.get_site_path("private", "files"))
		self._tmp.mkdir(parents=True, exist_ok=True)

	def _ok_path(self) -> Path:
		path = self._tmp / "liga_excel_ok_test.xlsx"
		_write_sample_xlsx(
			path,
			[
				_HEADERS,
				[
					"06/09/2026",
					"20:00",
					"22:00",
					CANCHA_3,
					"U15",
					"AZUL",
					"Rival A",
				],
				[
					"2026-09-07",
					"18:30",
					"20:00",
					"GIMNASIO 2",
					"Sub 13",
					"Nivel D",
					"Rival B",
				],
			],
		)
		return path

	def _file_url_for(self, path: Path, owner: str) -> str:
		doc = save_file(path.name, path.read_bytes(), None, None, is_private=1)
		frappe.db.set_value("File", doc.name, "owner", owner, update_modified=False)
		return doc.file_url

	def test_preview_sin_side_effects(self) -> None:
		path = self._ok_path()
		before = frappe.db.count("Reserva Espacio")
		preview = preview_excel_fixtures(path)
		self.assertEqual(len(preview["filas_ok"]), 2)
		self.assertEqual(preview["errores"], [])
		self.assertTrue(all(row["localia"] == "local" for row in preview["filas_ok"]))
		self.assertEqual(frappe.db.count("Reserva Espacio"), before)

	def test_apply_idempotente(self) -> None:
		path = self._ok_path()
		preview = preview_excel_fixtures(path)
		external_ids = [row["external_id"] for row in preview["filas_ok"]]
		self.assertEqual(len(set(external_ids)), 2)
		self.assertTrue(all(value.startswith("xlsx-") for value in external_ids))

		first = apply_excel_fixtures(path)
		self.assertEqual(first["creados"], 2)
		self.assertTrue(find_reserva_by_fixture(ORIGIN_LIGA_EXCEL, external_ids[0]))
		self.assertTrue(find_reserva_by_fixture(ORIGIN_LIGA_EXCEL, external_ids[1]))
		doc = frappe.get_doc(
			"Reserva Espacio",
			find_reserva_by_fixture(ORIGIN_LIGA_EXCEL, external_ids[1]),
		)
		self.assertEqual(doc.espacio, CANCHA_2)

		second = apply_excel_fixtures(path)
		self.assertEqual(second["creados"], 0)
		self.assertGreaterEqual(second["actualizados"], 2)

	def test_identidad_no_cambia_al_corregir_horario_o_espacio(self) -> None:
		base = {
			"fecha": "2026-09-07",
			"hora": "18:30",
			"espacio": CANCHA_2,
			"categoria": "Sub 13",
			"tira": "Nivel D",
			"rival": "Rival B",
		}
		changed = {**base, "hora": "19:00", "espacio": CANCHA_3}
		self.assertEqual(build_excel_external_id(base), build_excel_external_id(changed))

	def test_rechaza_extension_xls(self) -> None:
		path = self._tmp / "fixture_legacy.xls"
		path.write_bytes(b"legacy")
		with self.assertRaises(frappe.ValidationError):
			preview_excel_fixtures(path)

	def test_fila_invalida_no_aborta_lote(self) -> None:
		path = self._tmp / "liga_excel_mixed_test.xlsx"
		_write_sample_xlsx(
			path,
			[
				_HEADERS,
				[
					"06/09/2026",
					"20:00",
					"22:00",
					CANCHA_3,
					"U15",
					"AZUL",
					"Ok",
				],
				[
					"fecha-mala",
					"20:00",
					"22:00",
					CANCHA_3,
					"U15",
					"AZUL",
					"Bad",
				],
				[
					"08/09/2026",
					"19:00",
					"",
					"",
					"Superior",
					"Septima",
					"Sin espacio",
				],
			],
		)
		preview = preview_excel_fixtures(path)
		ok_id = preview["filas_ok"][0]["external_id"]
		result = apply_excel_fixtures(path)
		self.assertEqual(result["creados"], 1)
		self.assertTrue(find_reserva_by_fixture(ORIGIN_LIGA_EXCEL, ok_id))
		self.assertGreaterEqual(len(result["omitidos"]), 2)

	def test_whitelist_coordinacion_preview_apply(self) -> None:
		from club_management.spaces.api.fixtures_desk import (
			apply_fixtures_excel,
			preview_fixtures_excel,
		)

		path = self._ok_path()
		user = make_coordinacion_user("coord.excel@example.com")
		file_url = self._file_url_for(path, user)
		frappe.set_user(user)
		try:
			preview = preview_fixtures_excel(file_url)
			self.assertEqual(len(preview["filas_ok"]), 2)
			applied = apply_fixtures_excel(file_url, cancel_missing=0)
			self.assertIn("creados", applied)
		finally:
			frappe.set_user("Administrator")

	def test_coordinacion_rechaza_ruta_local_arbitraria(self) -> None:
		from club_management.spaces.api.fixtures_desk import preview_fixtures_excel

		path = self._ok_path()
		user = make_coordinacion_user("coord.excel.security@example.com")
		frappe.set_user(user)
		try:
			with self.assertRaises(frappe.PermissionError):
				preview_fixtures_excel(str(path))
		finally:
			frappe.set_user("Administrator")

	def test_coordinacion_descarga_template_canonico(self) -> None:
		from club_management.spaces.api.fixtures_desk import download_fixtures_excel_template

		user = make_coordinacion_user("coord.excel.template@example.com")
		frappe.set_user(user)
		try:
			frappe.local.response = frappe._dict()
			download_fixtures_excel_template()
			self.assertEqual(frappe.response["type"], "binary")
			self.assertEqual(frappe.response["filename"], "plantilla_fixture_ligas.xlsx")
			wb = openpyxl.load_workbook(BytesIO(frappe.response["filecontent"]), read_only=True)
			ws = wb["Fixture"]
			headers = [cell.value for cell in next(ws.iter_rows(max_row=1))]
			self.assertEqual(headers, _HEADERS)
			self.assertGreaterEqual(ws.max_row, 2)
		finally:
			frappe.set_user("Administrator")

	def test_ui_import_excel_ofrece_descarga_template(self) -> None:
		path = Path(
			frappe.get_app_path(
				"club_management",
				"public",
				"js",
				"ocupacion_espacios_page.js",
			)
		)
		source = path.read_text(encoding="utf-8")
		self.assertIn("download_fixtures_excel_template", source)
		self.assertIn("Descargar plantilla Excel", source)
