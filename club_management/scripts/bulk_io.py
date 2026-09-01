"""Helpers compartidos de CSV/Excel e identificación de socios para cargas masivas."""

from __future__ import annotations

import csv
from datetime import date, datetime
from pathlib import Path
from typing import Any

import frappe
from frappe import _
from frappe.utils import flt, getdate

SOCIO_DOCTYPE = "Socio"
PROD_SITE_MARK = "icdpedroechague.com.ar"
CONFIRM_LOCAL = "local-dev"
CONFIRM_PROD = "APPLY_PROD"


def is_production_site(site: str | None = None) -> bool:
	"""True si el site actual es producción ICDPE."""
	name = site if site is not None else str(getattr(frappe.local, "site", "") or "")
	return PROD_SITE_MARK in name


def ensure_bulk_apply_allowed(*, dry_run: bool, confirm: str = "") -> None:
	"""Gate apply masivo: local → `local-dev`; prod → `APPLY_PROD`. Dry-run siempre OK."""
	if dry_run:
		return
	token = (confirm or "").strip()
	if is_production_site():
		if token != CONFIRM_PROD:
			frappe.throw(
				_(
					"Apply bloqueado en {0}. Para producción: confirm='{1}'."
				).format(frappe.local.site, CONFIRM_PROD),
				frappe.ValidationError,
			)
		return
	if token != CONFIRM_LOCAL:
		frappe.throw(
			_("Apply bloqueado en {0}. Para local: confirm='{1}'.").format(
				frappe.local.site, CONFIRM_LOCAL
			),
			frappe.ValidationError,
		)


def ensure_not_production(*, dry_run: bool, confirm: str = "") -> None:
	"""Alias de `ensure_bulk_apply_allowed` (cargas masivas / ops)."""
	ensure_bulk_apply_allowed(dry_run=dry_run, confirm=confirm)


def cell(row: dict[str, Any], *keys: str) -> str:
	lookup = {(k or "").strip().lower(): (v if v is not None else "") for k, v in row.items()}
	for key in keys:
		if key.lower() in lookup:
			return str(lookup[key.lower()]).strip()
	return ""


def parse_monto(raw: str) -> float:
	text = (raw or "").strip().replace(" ", "").replace("$", "")
	if not text:
		return 0.0
	if "," in text and "." in text:
		if text.rfind(",") > text.rfind("."):
			text = text.replace(".", "").replace(",", ".")
		else:
			text = text.replace(",", "")
	elif "," in text:
		text = text.replace(".", "").replace(",", ".")
	return flt(text)


def parse_fecha(raw: str) -> date | None:
	text = (raw or "").strip()
	if not text:
		return None
	for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y"):
		try:
			return datetime.strptime(text, fmt).date()
		except ValueError:
			continue
	try:
		return getdate(text)
	except Exception:
		return None


def find_socio(*, nro_socio: str = "", dni: str = "") -> str | None:
	nro = (nro_socio or "").strip()
	if nro:
		if frappe.db.exists(SOCIO_DOCTYPE, nro):
			return nro
		meta = frappe.get_meta(SOCIO_DOCTYPE)
		if meta.has_field("numero_socio"):
			try:
				found = frappe.db.get_value(SOCIO_DOCTYPE, {"numero_socio": int(nro)}, "name")
			except ValueError:
				found = frappe.db.get_value(SOCIO_DOCTYPE, {"numero_socio": nro}, "name")
			if found:
				return found
		if meta.has_field("nro_socio_padron"):
			found = frappe.db.get_value(SOCIO_DOCTYPE, {"nro_socio_padron": nro}, "name")
			if found:
				return found
	doc = (dni or "").strip()
	if doc:
		found = frappe.db.get_value(SOCIO_DOCTYPE, {"dni": doc}, "name")
		if found:
			return found
	return None


def _read_csv(path: Path) -> list[dict[str, str]]:
	with path.open(encoding="utf-8-sig", newline="") as handle:
		return [{str(k): ("" if v is None else str(v)) for k, v in row.items()} for row in csv.DictReader(handle)]


def _xlsx_all_tuples(path: Path) -> list[tuple[Any, ...]]:
	from openpyxl import load_workbook

	wb = load_workbook(path, read_only=True, data_only=True)
	ws = wb.active
	rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
	wb.close()
	return rows


def _read_xlsx(path: Path) -> list[dict[str, str]]:
	try:
		from openpyxl import load_workbook  # noqa: F401
	except ImportError as exc:
		raise frappe.ValidationError(
			_("Para Excel (.xlsx) instalá openpyxl o exportá el archivo a CSV.")
		) from exc

	from club_management.scripts.excel_inscripciones import filas_from_socios_excel
	from club_management.scripts.informe_cobranzas import is_informe_cobranzas_header, pagos_from_informe_rows

	tuples = _xlsx_all_tuples(path)
	if not tuples:
		return []
	first = tuples[0][0] if tuples[0] else None
	if is_informe_cobranzas_header(first):
		return pagos_from_informe_rows(tuples)
	headers = [_norm_header(h) for h in tuples[0]]
	if "número de socio" in {h.lower() for h in headers} or "numero de socio" in {h.lower() for h in headers}:
		if any(h.lower() == "actividad" for h in headers):
			return filas_from_socios_excel(tuples)

	header_row = [str(h or "").strip() for h in tuples[0]]
	out: list[dict[str, str]] = []
	for values in tuples[1:]:
		row = {
			header_row[i]: ("" if (i >= len(values) or values[i] is None) else str(values[i]).strip())
			for i in range(len(header_row))
		}
		if any(row.values()):
			out.append(row)
	return out


def _norm_header(value: Any) -> str:
	return str(value or "").strip()


def read_bulk_rows(csv_path: str) -> list[dict[str, str]]:
	path = Path(csv_path)
	if not path.is_file():
		frappe.throw(_("Archivo no encontrado: {0}").format(csv_path), frappe.DoesNotExistError)
	if path.suffix.lower() in {".xlsx", ".xlsm"}:
		return _read_xlsx(path)
	return _read_csv(path)
